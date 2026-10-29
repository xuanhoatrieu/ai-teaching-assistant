import os
import sys
import struct
import wave
import zipfile
import tempfile
import unittest
from lxml import etree

sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx_service import PPTXGeneratorService
from main import _extract_slides_from_pptx, _render_pptx_slides_to_png

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
P = f"{{{NS_P}}}"


def run_gui_fix_slide_check(xml_bytes: bytes) -> bool:
    """
    Exact structural check mirroring fix_slide() from auto_audio_converter_gui.py.
    Returns True if any autoplay requirement was missing (meaning fix_slide would have had to patch it),
    or False if the XML is already 100% PowerPoint 'Start: Automatically' compliant.
    """
    root = etree.fromstring(xml_bytes)
    changed = False

    mseq = next((e for e in root.findall(f".//{P}cTn") if e.get("nodeType") == "mainSeq"), None)
    if mseq is None:
        return True

    mid = mseq.get("id", "2")
    seq = mseq.getparent()
    clist = mseq.find(f"{P}childTnLst")
    if clist is None:
        return True

    found_mediacall = False
    for par in list(clist):
        oc = par.find(f"{P}cTn")
        if oc is None:
            continue

        ac = None
        spid = None
        for c in oc.findall(f".//{P}cTn"):
            if c.get("presetClass") == "mediacall":
                cmd = c.find(f".//{P}cmd")
                if cmd is not None and "playFrom" in cmd.get("cmd", ""):
                    ac = c
                    s = c.find(f".//{P}spTgt")
                    if s is not None:
                        spid = s.get("spid")
                    break
        if ac is None:
            continue

        found_mediacall = True
        if ac.get("nodeType") != "afterEffect":
            changed = True

        sc = oc.find(f"{P}stCondLst")
        if sc is None:
            changed = True
        else:
            has_on_begin = any(
                x.get("evt") == "onBegin"
                and any(tn.get("val") == mid for tn in x.findall(f"{P}tn"))
                for x in sc.findall(f"{P}cond")
            )
            has_indefinite = any(x.get("delay") == "indefinite" for x in sc.findall(f"{P}cond"))
            if not (has_on_begin and has_indefinite):
                changed = True

        if seq is None or seq.find(f"{P}prevCondLst") is None or seq.find(f"{P}nextCondLst") is None:
            changed = True

        tmr = next((e for e in root.findall(f".//{P}cTn") if e.get("nodeType") == "tmRoot"), None)
        if tmr is None or not spid:
            changed = True
        else:
            tl = tmr.find(f"{P}childTnLst")
            has_bg_audio = tl is not None and any(
                s.get("spid") == spid
                for a in tl.findall(f"{P}audio")
                for s in a.findall(f".//{P}spTgt")
            )
            if not has_bg_audio:
                changed = True

    return (not found_mediacall) or changed


class TestPptxAutoplayAndFullSlideRender(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.wav_path = os.path.join(self.tmp_dir.name, "sample.wav")
        with wave.open(self.wav_path, "wb") as wf:
            wf.setnchannels(1)
            wf.setsampwidth(2)
            wf.setframerate(24000)
            wf.writeframes(struct.pack("<100h", *([0] * 100)))
        self.service = PPTXGeneratorService()

    def tearDown(self):
        self.tmp_dir.cleanup()

    def test_fresh_slide_autoplay_matches_powerpoint_and_ispring_ooxml(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # Append an <p:extLst> to verify <p:timing> is inserted BEFORE <p:extLst> per CT_Slide XSD
        slide._element.append(parse_xml(f'<p:extLst xmlns:p="{NS_P}"/>'))

        self.service._add_audio_with_autoplay(slide, self.wav_path)

        out_pptx = os.path.join(self.tmp_dir.name, "fresh.pptx")
        prs.save(out_pptx)

        with zipfile.ZipFile(out_pptx, "r") as zf:
            slide_xml = zf.read("ppt/slides/slide1.xml")

        # 1. Must require ZERO patching by auto_audio_converter_gui.py fix_slide logic
        self.assertFalse(
            run_gui_fix_slide_check(slide_xml),
            "Generated slide XML should already satisfy 100% of PowerPoint Start: Automatically OOXML rules",
        )

        # 2. Verify <p:timing> precedes <p:extLst>
        root = etree.fromstring(slide_xml)
        child_tags = [etree.QName(el).localname for el in root]
        self.assertIn("timing", child_tags)
        self.assertIn("extLst", child_tags)
        self.assertLess(child_tags.index("timing"), child_tags.index("extLst"))

    def test_uploaded_slide_preserves_existing_animations_when_injecting_autoplay_audio(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        tb.text_frame.text = "Animated Heading"
        anim_spid = tb.shape_id

        # Simulate pre-existing user entrance animation inside <p:timing>
        pre_existing_timing = f'''
        <p:timing xmlns:p="{NS_P}">
            <p:tnLst>
                <p:par>
                    <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
                        <p:childTnLst>
                            <p:seq concurrent="1" nextAc="seek">
                                <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                                    <p:childTnLst>
                                        <p:par>
                                            <p:cTn id="10" fill="hold">
                                                <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
                                                <p:childTnLst>
                                                    <p:par>
                                                        <p:cTn id="11" presetID="10" presetClass="entr" fill="hold" nodeType="clickEffect">
                                                            <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                                            <p:childTnLst>
                                                                <p:set>
                                                                    <p:cBhvr>
                                                                        <p:cTn id="12" dur="1" fill="hold"/>
                                                                        <p:tgtEl><p:spTgt spid="{anim_spid}"/></p:tgtEl>
                                                                    </p:cBhvr>
                                                                </p:set>
                                                            </p:childTnLst>
                                                        </p:cTn>
                                                    </p:par>
                                                </p:childTnLst>
                                            </p:cTn>
                                        </p:par>
                                    </p:childTnLst>
                                </p:cTn>
                            </p:seq>
                        </p:childTnLst>
                    </p:cTn>
                </p:par>
            </p:tnLst>
        </p:timing>
        '''
        slide._element.append(parse_xml(pre_existing_timing))

        self.service._add_audio_with_autoplay(slide, self.wav_path)

        out_pptx = os.path.join(self.tmp_dir.name, "animated.pptx")
        prs.save(out_pptx)

        with zipfile.ZipFile(out_pptx, "r") as zf:
            slide_xml = zf.read("ppt/slides/slide1.xml")

        self.assertFalse(run_gui_fix_slide_check(slide_xml))
        root = etree.fromstring(slide_xml)
        # Ensure original entrance animation (id="11", presetClass="entr") is still intact
        entr_nodes = [
            e for e in root.findall(f".//{P}cTn")
            if e.get("id") == "11" and e.get("presetClass") == "entr"
        ]
        self.assertEqual(len(entr_nodes), 1)
        # Ensure all cTn ids are unique
        ids = [e.get("id") for e in root.findall(f".//{P}cTn") if e.get("id")]
        self.assertEqual(len(ids), len(set(ids)), f"Duplicate cTn IDs found: {ids}")

    def test_extract_filters_page_numbers_and_renders_full_slide_pngs(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        for idx, heading in enumerate(["Bài 1: Ngữ âm /iː/ & /ɪ/", "Bảng so sánh Minimal Pairs"]):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(10), Inches(1))
            tbox.text_frame.text = heading

            cbox = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(10), Inches(2))
            cbox.text_frame.text = "Phát âm chuẩn tiếng Anh\x0b(Ví dụ: sheep /ʃiːp/ vs ship /ʃɪp/)"

            # Standalone page-number badge that used to pollute bullets
            pnum = slide.shapes.add_textbox(Inches(12.2), Inches(6.8), Inches(0.6), Inches(0.4))
            pnum.text_frame.text = str(idx + 1)

        src_pptx = os.path.join(self.tmp_dir.name, "lecture.pptx")
        prs.save(src_pptx)

        parsed = _extract_slides_from_pptx(Presentation(src_pptx))
        self.assertEqual(len(parsed), 2)
        self.assertEqual(parsed[0]["title"], "Bài 1: Ngữ âm /iː/ & /ɪ/")
        self.assertNotIn("1", parsed[0]["content"])
        self.assertNotIn("\x0b", parsed[0]["content"][0])

        rendered_dir = os.path.join(self.tmp_dir.name, "rendered_slides")
        rendered_map = _render_pptx_slides_to_png(src_pptx, rendered_dir, expected_count=2, dpi=120)
        self.assertEqual(len(rendered_map), 2)
        for i in (0, 1):
            png_path = os.path.join(rendered_dir, rendered_map[i])
            self.assertTrue(os.path.isfile(png_path))
            with open(png_path, "rb") as f:
                header = f.read(8)
            self.assertEqual(header, b"\x89PNG\r\n\x1a\n")


if __name__ == "__main__":
    unittest.main()
